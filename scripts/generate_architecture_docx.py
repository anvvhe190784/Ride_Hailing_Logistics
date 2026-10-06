# -*- coding: utf-8 -*-
"""
Script to generate PROJECT_ARCHITECTURE_AND_DIAGRAMS.docx
Creates a beautifully formatted Word document covering all system architecture,
features, and diagrams.
"""
import os
import re
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT_DOCX = r"d:\Rikkei\Rikkei FE\Ride_Hailing_Logistics\docs\PROJECT_ARCHITECTURE_AND_DIAGRAMS.docx"
SOURCE_MD = r"d:\Rikkei\Rikkei FE\Ride_Hailing_Logistics\docs\PROJECT_ARCHITECTURE_AND_DIAGRAMS.md"

W_VAL = 'w:val'
W_COLOR = 'w:color'

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn(W_VAL), 'clear')
    shd.set(qn(W_COLOR), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def add_code_box(doc, code_text, title="MERMAID SƠ ĐỒ"):
    table = doc.add_table(1, 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, 'F1F5F9')
    
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left')
    left.set(qn(W_VAL), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn(W_COLOR), '2563EB')
    tc_borders.append(left)
    for s in ('top', 'bottom', 'right'):
        b = OxmlElement(f'w:{s}')
        b.set(qn(W_VAL), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn(W_COLOR), 'CBD5E1')
        tc_borders.append(b)
    tc_pr.append(tc_borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📊 {title}\n")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(37, 99, 235)
    
    p_code = cell.add_paragraph()
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.space_after = Pt(3)
    r_code = p_code.add_run(code_text.strip())
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

def re_hyphen(s):
    return set(s.strip()).issubset({'-', ':'})

def add_formatted_text(paragraph, text):
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]*`)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(16, 75, 120)
        else:
            r = paragraph.add_run(token)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(51, 65, 85)

def format_table_headers(table, headers):
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, '0F4C81')
        p_cell = cell.paragraphs[0]
        p_cell.paragraph_format.space_before = Pt(3)
        p_cell.paragraph_format.space_after = Pt(3)
        r = p_cell.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

def format_table_rows(table, rows_data, num_cols):
    for row_idx, row in enumerate(rows_data, start=1):
        bg_color = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
        for col_idx, text in enumerate(row):
            if col_idx < num_cols:
                cell = table.cell(row_idx, col_idx)
                set_cell_background(cell, bg_color)
                p_cell = cell.paragraphs[0]
                p_cell.paragraph_format.space_before = Pt(2)
                p_cell.paragraph_format.space_after = Pt(2)
                r = p_cell.add_run(text)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

def parse_markdown_table(doc, lines, current_idx):
    table_lines = []
    i = current_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        table_lines.append(lines[i].strip())
        i += 1

    rows_data = []
    for t_line in table_lines:
        parts = [p.strip() for p in t_line.split('|')[1:-1]]
        if not all(re_hyphen(p) for p in parts):
            rows_data.append(parts)

    if rows_data:
        num_cols = len(rows_data[0])
        table = doc.add_table(len(rows_data), num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        format_table_headers(table, rows_data[0])
        format_table_rows(table, rows_data[1:], num_cols)
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(4)

    return i

def add_heading(doc, text, level):
    h = doc.add_paragraph()
    if level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(15, 76, 129)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(30, 64, 175)

def render_content_line(doc, line):
    stripped = line.strip()
    if not stripped or line.startswith('# '):
        return
    if line.startswith('## '):
        add_heading(doc, line[3:].strip(), 2)
    elif line.startswith('### '):
        add_heading(doc, line[4:].strip(), 3)
    elif stripped.startswith(('- ', '* ')):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        add_formatted_text(p, stripped[2:].strip())
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_text(p, stripped)

def setup_document():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_t = title_p.add_run("TÀI LIỆU TOÀN DIỆN VỀ KIẾN TRÚC, CHỨC NĂNG & DIAGRAMS\nHỆ THỐNG RIDE-HAILING & LOGISTICS (SRS-RHL-002)")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(16)
    run_t.font.color.rgb = RGBColor(15, 76, 129)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Phiên bản: 1.0 | Nền tảng: Java 21, Spring Boot 3.4.4, Spring Cloud, PostgreSQL 18 PostGIS, Redis 7, Kafka 3")
    run_sub.italic = True
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    return doc

def main():
    print(f"Reading source markdown from: {SOURCE_MD}")
    with open(SOURCE_MD, 'r', encoding='utf-8') as f:
        md_content = f.read()

    doc = setup_document()
    lines = md_content.split('\n')
    i = 0
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code_block:
                title = f"SƠ ĐỒ TRỰC QUAN ({code_block_lang.upper()})" if code_block_lang else "ĐOẠN MÃ LỆNH / SCHEMA"
                add_code_box(doc, '\n'.join(code_block_lines), title)
                in_code_block = False
            else:
                in_code_block = True
                code_block_lang = line.strip()[3:].strip()
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            i = parse_markdown_table(doc, lines, i)
            continue

        render_content_line(doc, line)
        i += 1

    print(f"Saving formatted architecture document to: {OUTPUT_DOCX}")
    os.makedirs(os.path.dirname(OUTPUT_DOCX), exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print("Architecture DOCX successfully generated!")

if __name__ == "__main__":
    main()
