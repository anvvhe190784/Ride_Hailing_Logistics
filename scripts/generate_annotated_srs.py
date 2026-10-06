# -*- coding: utf-8 -*-
import os
import re
import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Import the 253 requirement annotations
from requirement_annotations_data import ANNOTATIONS_DB

SRC_DOCX = r"C:\Users\vanan\.gemini\antigravity\brain\0933d14e-bd9f-424f-a8e9-adb4fcce4a4b\.user_uploaded\media_1791265125716.docx"
DST_DOCX = r"d:\Rikkei\Rikkei FE\Ride_Hailing_Logistics\docs\SRS_RHL_002_Annotated_Implementation.docx"

W_VAL = 'w:val'
W_COLOR = 'w:color'

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn(W_VAL), 'clear')
    shd.set(qn(W_COLOR), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def create_callout_box(doc, target_p, req_code, solution, code_loc, db_loc):
    table = doc.add_table(1, 1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, 'F4F7FB')
    
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    
    # Left border: thick navy/blue (3pt)
    left = OxmlElement('w:left')
    left.set(qn(W_VAL), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '0')
    left.set(qn(W_COLOR), '0F4C81')
    tc_borders.append(left)
    
    # Top/Right/Bottom: subtle border (0.5pt)
    for side in ('top', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn(W_VAL), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn(W_COLOR), 'D1DCE5')
        tc_borders.append(b)
    tc_pr.append(tc_borders)
    
    # Title paragraph
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f'📌 GIẢI PHÁP THỰC THI & VỊ TRÍ MÃ NGUỒN [{req_code}]\n')
    run_title.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.name = 'Arial'
    run_title.font.color.rgb = RGBColor(15, 76, 129)
    
    # Details paragraph
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    
    # Solution
    r1 = p2.add_run('⚙️ Cách giải quyết: ')
    r1.bold = True
    r1.font.size = Pt(9)
    r1.font.name = 'Arial'
    r1.font.color.rgb = RGBColor(30, 41, 59)
    r2 = p2.add_run(solution + '\n')
    r2.font.size = Pt(9)
    r2.font.name = 'Arial'
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    # Code location
    r3 = p2.add_run('📁 Vị trí Code & Module: ')
    r3.bold = True
    r3.font.size = Pt(9)
    r3.font.name = 'Arial'
    r3.font.color.rgb = RGBColor(30, 41, 59)
    r4 = p2.add_run(code_loc + '\n')
    r4.font.size = Pt(8.5)
    r4.font.name = 'Consolas'
    r4.font.color.rgb = RGBColor(16, 75, 120)
    
    # DB & Messaging
    r5 = p2.add_run('🗄️ Database & Messaging: ')
    r5.bold = True
    r5.font.size = Pt(9)
    r5.font.name = 'Arial'
    r5.font.color.rgb = RGBColor(30, 41, 59)
    r6 = p2.add_run(db_loc)
    r6.font.size = Pt(8.5)
    r6.font.name = 'Consolas'
    r6.font.color.rgb = RGBColor(39, 103, 73)
    
    # Relocate table directly after target paragraph
    target_p._p.addnext(table._tbl)

def main():
    print(f"Loading source SRS document from: {SRC_DOCX}")
    doc = docx.Document(SRC_DOCX)
    print(f"Document loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")
    
    # Regex to match requirement prefixes
    pattern = re.compile(r'^(CON-\d+|BR-\d+|[A-Z0-9]+-[A-Z0-9]+-\d+):\s*(.*)')
    
    annotated_count = 0
    matched_reqs = []
    
    # Identify target paragraphs first
    targets = []
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        m = pattern.match(text)
        if m:
            req_code = m.group(1)
            targets.append((p, req_code))
            
    print(f"Found {len(targets)} requirement paragraphs to annotate.")
    
    # Process annotations
    for p, req_code in targets:
        if req_code in ANNOTATIONS_DB:
            info = ANNOTATIONS_DB[req_code]
            create_callout_box(doc, p, req_code, info["solution"], info["code"], info["db"])
            annotated_count += 1
            matched_reqs.append(req_code)
        else:
            print(f"Warning: No annotation found for {req_code}")
            
    print(f"Successfully created callout boxes for {annotated_count} requirements.")
    
    # Ensure target directory exists
    os.makedirs(os.path.dirname(DST_DOCX), exist_ok=True)
    
    print(f"Saving annotated document to: {DST_DOCX}")
    doc.save(DST_DOCX)
    print("Done! Document saved successfully.")

if __name__ == "__main__":
    main()
